# -*- coding: utf-8 -*-
"""Tao QS_DAM_NhapLieu.xlsx (V2) - bo cuc giong sheet DCE_Pro_Beam (DAM.xlsm), toi uu + mo rong.

Vi tri o NHAP LIEU giu dung nhu DCE_Pro_Beam (F2:F8, J2:J8, N2:N8, B11, B12, C11:AG30) de
QS_DAM.lsp doc bang bo doc DCE; phan mo rong: thep cho S3:T7, tai san / san lat Y2:Y3.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "QS_DAM_NhapLieu.xlsx"

NSPAN = 15                       # 15 nhip / 16 goi : cot C .. AG (giong DCE)
C0 = 3                           # cot C = Goi 1
CLAST = C0 + 2 * NSPAN           # cot AG = Goi 16
GCOLS = [C0 + 2 * i for i in range(NSPAN + 1)]
NCOLS = [C0 + 2 * i + 1 for i in range(NSPAN)]
FIRST, LAST = CL(C0), CL(CLAST)

# ---------------------------------------------------------------- style
FONT = "Arial"
def F(**k):
    k.setdefault("name", FONT); k.setdefault("size", 10)
    return Font(**k)
def FILL(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="808080")
hair = Side(style="hair", color="A6A6A6")
med = Side(style="medium", color="000000")
B_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
B_HAIR = Border(left=thin, right=thin, top=hair, bottom=hair)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
NOWRAP = Alignment(horizontal="left", vertical="center", wrap_text=False)

C_TITLE = "FFC000"   # cam DCE
C_LABEL = "F4F4AC"   # vang nhat DCE (nhan / o thep)
C_LAYERG = "E6E49A"  # o thep tai cot goi (dam hon mot chut)
C_GOIHDR = "D9D9D9"
C_NHIPHDR = "DDEBF7"
C_GOIDARK = "808080" # cot goi dong 23..30 (giong DCE)
C_NA = "595959"
C_AUTH = "92D050"
C_CALC = "EDEDED"
C_SEC = "BDD7EE"

# ---------------------------------------------------------------- danh sach chon
LISTS = {
    "THEP": ["-"] + [f"{n}t{d}" for d in (14, 16, 18, 20, 22, 25, 28, 32) for n in (2, 3, 4, 5, 6)],
    "DAI": ["8-100/200", "8-150/200", "10-100/150", "10-100/200", "10-150/200", "12-100/150",
            "12-100/200", "12-150/200", "8-100", "10-100", "12-100"],
    "GIA": ["2x2t12", "2x2t14", "2x2t16", "3x2t12", "3x2t14", "2t12", "2t14", "4t12"],
    "DAITRONG": ["8-100/200", "10-100/200", "10-150/200", "T", "B", "T/B", "TL", "TR", "BL", "BR",
                 "8-100/200/T", "8-100/200/B"],
    "TLINK": ["T.LINK"],
    "DAIGC": ["10t10a50", "8t8a50", "8t10a50", "10t10a50L", "10t10a50R", "8u8a50"],
    "VAIBO": ["2t16", "2t18", "2t20", "2t22"],
    "ANCOT": ["1", "2", "3"],
    "DOITHEP": ["0;3t25", "0;3t28", "0;5t28", "0;25", "0;28"],
    "KGOI": ["0.25|0.25|0.15", "0.3|0.25|0.15", "0.25/0.25/0.15", "0.3/0.25/0.15", "0.25\\0.25\\0.15"],
    "KNHIP": ["0.15\\0.25\\0.1", "0.15|0.25|0.1", "0.15/0.25/0.1"],
    "NEO": ["40", "10;12;14;16-30d/40d", "10-500mm/12-600mm/14-690mm/16-790mm/18-890mm/20-990mm/22-1090mm/25-1230mm/28-1380mm/32-1580mm"],
    "DAIC": ["None", "So le", "So le 2", "Toàn bộ"],
    "GIATCD": ["Mép trên", "Mép dưới", "Tùy chỉnh"],
    "CDOKC": ["-1500", "1500", "1500mm", "0"],
    "KCDAIC": ["400", "400mm", "300/600"],
    "LECHMCN": ["Ltt", "0"],
    "NCK": ["1", "2", "1/T", "1/B", "1/T/B"],
    "CHOKIEU": ["Khong", "Cho thang", "Coupler"],
    "CHOLOP": ["TBG", "TB", "T", "B"],
    "CHOTHEP": ["Tat ca", "Chay suot"],
    "CHOL": ["AUTO", "40d", "50d", "1200", "0", "100", "100/300"],
    "KHONGCO": ["Khong", "Co"],
    "DAICON": ["AUTO", "0", "2", "3", "2_3", "2_4", "2-4", "3,2_4", "2,4"],
    "TAISAN": ["2 ben", "Trai", "Phai", "Khong"],
}
LISTCOL = {k: CL(i + 1) for i, k in enumerate(LISTS)}

def list_ref(key):
    c = LISTCOL[key]
    return f"LIST!${c}$2:${c}${len(LISTS[key]) + 1}"

# ---------------------------------------------------------------- du lieu vi du (dam L1.B137 - DAM.xlsm)
EX_HEAD = {
    "F2": "L1.B137-1,2,3,4,5,6,chờ 2,6", "F3": 400, "F4": 1000, "F5": "1", "F6": -0.1, "F7": 200, "F8": 2.5,
    "J2": "0.25|0.25|0.15", "J3": "0.15\\0.25\\0.1",
    "J4": "10-500mm/12-600mm/14-690mm/16-790mm/18-890mm/20-990mm/22-1090mm/25-1230mm/28-1380mm/32-1580mm",
    "J5": "10-350mm/12-420mm/14-480mm/16-550mm/18-620mm/20-690mm/22-760mm/25-860mm/28-960mm/32-1100mm",
    "J6": "10-350mm/12-420mm/14-480mm/16-550mm/18-620mm/20-690mm/22-760mm/25-860mm/28-960mm/32-1100mm",
    "N2": 32, "N3": "-1500", "N4": "So le 2", "N5": 10, "N6": "400", "N7": "Mép dưới",
    "B11": "3t28", "B12": "3t28",
    "S3": "Khong", "T3": "Khong", "S4": "TB", "T4": "TB", "S5": "Tat ca", "T5": "Tat ca",
    "S6": "AUTO", "T6": "AUTO", "S7": "Khong", "T7": "Khong", "S8": "", "T8": "", "Y2": "2 ben", "Y3": "Khong",
}
# luoi: {dong: [gia tri theo cot C, D, E ...]}
EX_GRID = {
    11: [300, 7250, 900, 8200, 900, 7500, 900, 7550, 800, 6940, 760, 8700, 1000],
    12: [None, None, None, None, None, 500, None, None, None, 400],
    14: ["3t28", None, "3t28", None, "3t28;5t28", None, "5t28", None, "5t28;3t28", None, "3t28", None, "-"],
    15: ["2t25", None, "3t28", None, "2t25", None, "5t28", None, ";2t28", None, None, None, "2t25"],
    16: [None, None, None, None, None, None, "3t28"],
    20: [None, "3t25", None, "2t25"],
    21: [None, "3t28", None, "3t28", None, "5t28", None, "3t28", None, "2t25", None, "2t25"],
    23: [None, None, None, None, None, "0;5t28", None, None, None, "0;3t28"],
    24: [None, "2x2t12", None, "2x2t12", None, "2x2t12", None, "2x2t12", None, "2x2t12", None, "2x2t12"],
    25: [None, None, None, None, None, "0;5t28", None, None, None, "0;3t28"],
    26: ["E", "10-100/150", "6", "10-100/150", "7", "12-100/150", "8", "12-100/150", "1", "10-100/200", "2",
         "10-100/200", "3"],
    27: [-149.2, None, 700, None, -0.8, None, 0, None, 0, None, -20.1, None, 0.3],
    28: ["400x600", None, "400x600", None, "400x600", None, "400x600", None, "300x600", None, "300x600"],
    29: [3950, "10t10a50", 4900, "10t10a50", 4175, "10t10a50", 3120, "10t10a50", 3300, "10t10a50", 2260,
         "10t10a50"],
}
EX_AS = {37: [35.2, None, "48.5/50.1", None, "42.0/65.0", None, "98.6/101.3", None, "55.0/45.0", None, "30.2/31.0", None, 20.5],
         40: [None, 45.3, None, 40.2, None, 58.1, None, 50.2, None, 25.0, None, 26.3]}

# ---------------------------------------------------------------- dong luoi: nhan + goi y theo goi / nhip
ROWS = {
    11: ("Thép chạy suốt TRÊN  |  Gối: bề rộng  ·  Nhịp: L thông thủy", None,
         ("Gối: bề rộng (mm)", "300 = cột 300\n0 = console (đầu tự do)\n220x500 = gối là DẦM 220x500\n1000/220 = cột dưới 1000, cột trên 220\n1000/200/-400 = cột trên 200 lệch tim -400", None),
         ("Nhịp: L thông thủy (mm)", "Khoảng cách 2 mép gối (mm).\nĐể trống = hết dầm (các nhịp phải liền nhau).\nDòng 34-35: suy Ltt từ khoảng cách trục.", None)),
    12: ("Thép chạy suốt DƯỚI  |  Gối: đai trong cột  ·  Nhịp: b nhịp", None,
         ("Gối: bước đai trong cột", "Bước đai trong vùng cột (mm), vd 200.\n+200 = ghi kèm (+).\nĐể trống = không vẽ đai trong cột.", None),
         ("Nhịp: bề rộng nhịp (mm)", "Bề rộng dầm tại nhịp nếu KHÁC b dầm (F3).\nĐể trống = b dầm.", None)),
    23: ("Gối: ẩn cột  ·  Nhịp: đổi thép chạy TRÊN", None,
         ("Gối: ẩn cột khi vẽ", "1 = ẩn cột phía TRÊN\n2 = ẩn cột phía DƯỚI\n3 = ẩn cả hai\nĐể trống = vẽ đủ.", "ANCOT"),
         ("Nhịp: đổi thép chạy TRÊN", "0;5t28 = từ nhịp này thép chạy TRÊN đổi thành 5t28\n0;25 = chỉ đổi đường kính\nSố đầu ≠ 0 (giật mép trên): QS_DAM chưa vẽ giật cấp, chỉ đổi thép.", "DOITHEP")),
    24: ("Gối: dầm giao tại cột  ·  Nhịp: thép giá", None,
         ("Gối: dầm giao tại cột", "KT dầm giao tại cột (vẽ nét khuất), vd 400x600.", None),
         ("Nhịp: thép giá", "2x2t12 = 2 lớp, mỗi lớp 2t12\n4t12 = hiểu 2 lớp x 2t12\nNhịp liền kề cùng thép giá được gộp 1 thanh.\nĐể trống = không có.", "GIA")),
    25: ("Gối: lệch dầm giao  ·  Nhịp: đổi thép chạy DƯỚI", None,
         ("Gối: lệch dầm giao / tim cột", "Độ lệch tim dầm giao tại cột so với tim cột (mm), vd -150.", None),
         ("Nhịp: đổi thép chạy DƯỚI", "0;5t28 = từ nhịp này thép chạy DƯỚI đổi thành 5t28\n0;25 = chỉ đổi đường kính\nSố đầu ≠ 0 (giật mép dưới): chưa vẽ giật cấp.", "DOITHEP")),
    26: ("Gối: tên trục  ·  Nhịp: thép đai", None,
         ("Gối: tên trục", "Tên trục tại gối, vd E hoặc 6.\nNhiều dòng (Alt+Enter): lấy dòng cuối.", None),
         ("Nhịp: thép đai", "Ø-bước gối/bước giữa nhịp, vd 10-100/150\n10-100 = rải đều.\nVùng đai dày = hệ số thứ 2 ô J2.\nNhịp Ltt < F8 (m): rải đều bước gối.", "DAI")),
    27: ("Gối: lệch trục  ·  Nhịp: đai trong / cắt thép", None,
         ("Gối: lệch trục so tim gối", "Độ lệch trục so với tim gối (mm), vd -149.2.\nDương = trục lệch sang phải.", None),
         ("Nhịp: đai trong / cắt thép", "6-100/200 = đai trong (đai con)\nT / B / T/B = mối nối bắt buộc thép chạy TRÊN / DƯỚI\nTL TR BL BR = nối phía trái / phải nhịp\nGhép: 6-100/200/B", "DAITRONG")),
    28: ("Gối: dầm giao trong nhịp phải  ·  Nhịp: kiểu đai", None,
         ("Gối: dầm giao ở nhịp phải", "KT dầm giao nằm trong NHỊP BÊN PHẢI gối này, vd 400x600.\nNhiều dầm: 400x600/300x600", None),
         ("Nhịp: kiểu đai", "T.LINK = đai xoắn (U bao + C mũ).\nĐể trống = đai kín.", "TLINK")),
    29: ("Gối: vị trí dầm giao  ·  Nhịp: đai gia cường", None,
         ("Gối: vị trí dầm giao (mm)", "Khoảng cách từ TIM gối này tới tim dầm giao trong nhịp phải, vd 3950.\nNhiều dầm: 2000/5000", None),
         ("Nhịp: đai gia cường", "10t10a50 = 10 đai Ø10 bước 50, chia 2 bên dầm giao\n10t10a50L / R = chỉ bên trái / phải\nNhiều dầm: 10t10a50/8t8a50", "DAIGC")),
    30: ("Nhịp: thép vai bò tại dầm giao", None,
         None,
         ("Nhịp: thép vai bò", "Thép vai bò tại dầm giao trong nhịp, vd 2t18.", "VAIBO")),
}
for k in range(1, 6):
    ROWS[12 + k] = (None, f"Lớp {k}",
                    (f"Gối: TC trên lớp {k}", "Thép mũ gối: 3t28 (2 bên như nhau)\n3t28;5t28 = trái;phải\n;2t28 = chỉ bên phải\n2t28+1t25 = nhiều loại, - = không có\nVươn theo hệ số 1 ô J2." + ("\nLớp 1 = cùng lớp thép chạy suốt." if k == 1 else ""), "THEP"),
                    (f"Nhịp: TC trên lớp {k}", "Thép tăng cường TRÊN ở bụng nhịp, vd 2t25.\nChiều dài theo hệ số 3 ô J2." + ("\nLớp 1 = cùng lớp thép chạy suốt." if k == 1 else ""), "THEP"))
    ROWS[23 - k] = (None, f"Lớp {k}",
                    (f"Gối: TC dưới lớp {k}", "Thép tăng cường DƯỚI tại gối, vd 2t25.\n3t25;2t25 = trái;phải\nVươn theo hệ số 2 ô J3." + ("\nLớp 1 = cùng lớp thép chạy suốt." if k == 1 else ""), "THEP"),
                    (f"Nhịp: TC dưới lớp {k}", "Thép tăng cường DƯỚI ở bụng nhịp, vd 3t28.\nCách mép gối hệ số 1 ô J3 (nhịp đầu/cuối: hệ số 3)." + ("\nLớp 1 = cùng lớp thép chạy suốt." if k == 1 else ""), "THEP"))

# ---------------------------------------------------------------- tro giup dai (comment)
HELP_ROW = {
    11: "DÒNG 11\n• Ô B11: thép chạy suốt TRÊN, vd 3t28 (2t28+1t25 nếu nhiều loại).\n• Cột GỐI: bề rộng gối (mm)\n  300 = cột 300 · 0 = console\n  220x500 = gối là dầm 220x500\n  1000/220 = cột dưới 1000, cột trên 220\n  1000/200/-400 = cột trên 200, tim lệch -400\n• Cột NHỊP: chiều dài THÔNG THỦY (mm). Nhịp phải liền nhau, ô trống = hết dầm.",
    12: "DÒNG 12\n• Ô B12: thép chạy suốt DƯỚI, vd 3t28.\n• Cột GỐI: bước đai trong cột (mm), vd 200 hoặc +200.\n• Cột NHỊP: bề rộng nhịp nếu khác b dầm.\n(Móng băng dạng 2t16/12a100/12a100: chưa hỗ trợ.)",
    13: "THÉP TĂNG CƯỜNG LỚP TRÊN (dòng 13 = lớp 1 ... 17 = lớp 5)\n• Cột GỐI: thép mũ gối. 3t28 = 2 bên như nhau; 3t28;5t28 = trái;phải; ;2t28 = chỉ bên phải.\n• Cột NHỊP: thép tăng cường bụng nhịp phía trên.\n• '-' = không có. Lớp 1 cùng cao độ thép chạy suốt.",
    18: "THÉP TĂNG CƯỜNG LỚP DƯỚI (dòng 22 = lớp 1 ... 18 = lớp 5, giống DCE)\n• Cột NHỊP: thép tăng cường bụng nhịp phía dưới.\n• Cột GỐI: thép tăng cường dưới tại gối.",
    23: "DÒNG 23\n• GỐI: ẩn cột khi vẽ (1 trên · 2 dưới · 3 cả hai).\n• NHỊP: giật mép trên + đổi thép chạy TRÊN từ nhịp này:\n  100 = mép trên giật lên 100 (chưa vẽ giật cấp)\n  0;25 = đổi đường kính thép chạy thành 25\n  0;3t25 = đổi thép chạy thành 3t25",
    24: "DÒNG 24\n• GỐI: kích thước dầm giao tại cột, vd 400x600.\n• NHỊP: thép giá. 2x2t12 = 2 lớp x 2t12; 4t12 = 2 lớp x 2t12.",
    25: "DÒNG 25\n• GỐI: độ lệch dầm giao tại cột so với tim cột (mm).\n• NHỊP: giật mép dưới + đổi thép chạy DƯỚI (cú pháp như dòng 23).",
    26: "DÒNG 26\n• GỐI: tên trục (nhiều dòng thì lấy dòng cuối).\n• NHỊP: thép đai Ø-bước gối/bước nhịp, vd 10-100/200; 10-100 = rải đều.",
    27: "DÒNG 27\n• GỐI: độ lệch trục so với tim gối (mm).\n• NHỊP: đai trong 6-100/200 (để trống: tự thêm khi lớp 1 trên ≥ 4 thanh). Kiểu đai trong chọn ở QS_DAMSET trang 2:\n  1 nhánh tại mỗi thanh giữa (như DCE) hoặc đai kín ôm thanh 2 và n-1.\n  và/hoặc mối nối bắt buộc: T, B, T/B, TL, TR, BL, BR. Ví dụ 6-100/200/B.",
    28: "DÒNG 28\n• GỐI: KT dầm giao nằm trong NHỊP BÊN PHẢI gối (400x600 hoặc 400x600/300x600).\n• NHỊP: T.LINK = đai xoắn (U bao + C mũ). Mã giật sàn DCE (T.LINK/-100/-200/1) được lưu, chưa vẽ.",
    29: "DÒNG 29\n• GỐI: khoảng cách từ TIM gối tới tim dầm giao trong nhịp phải (2000 hoặc 2000/5000).\n• NHỊP: đai gia cường tại dầm giao: 8t8a50 (chia 2 bên), 8t8a50L / 8t8a50R (1 bên), 8u8a50 (chữ u = đai U).\n  Mỗi đai gia cường có thêm bộ đai trong / đai 1 nhánh như đai thường (tắt được ở QS_DAMSET trang 2).",
    30: "DÒNG 30 (chỉ cột NHỊP): thép vai bò tại dầm giao, vd 2t16.",
}

HEAD = [  # (o nhan, merge, o gia tri, nhan, prompt ngan (<=255), list key, help dai, strict)
    ("C2", "C2:E2", "F2", "Tên dầm:", "Tên dầm (ghi trên bản vẽ, tên file CSV).", None, None, False),
    ("C3", "C3:E3", "F3", "b dầm (mm):", "Bề rộng dầm (mm), vd 400.", None, "b dầm (mm). Nhịp có b khác: nhập ở dòng 12 cột nhịp.\nMóng băng dạng 300/500/500: chưa hỗ trợ (lấy số đầu).", False),
    ("C4", "C4:E4", "F4", "h dầm (mm):", "Chiều cao dầm (mm), vd 1000.", None, None, False),
    ("C5", "C5:E5", "F5", "Số cấu kiện:", "1 = số cấu kiện giống nhau\n1/T = cắt toàn bộ thép chạy TRÊN\n1/T/B = cắt cả TRÊN và DƯỚI", "NCK",
     "Số cấu kiện giống nhau (nhân số lượng thống kê).\nThêm /T, /B: đặt mối nối bắt buộc cho thép chạy suốt TRÊN / DƯỚI ở mọi nhịp.\nVí dụ: 1/T ; 1/T/B", False),
    ("C6", "C6:E6", "F6", "Cao trình dầm (m):", "Cao độ mặt dầm (m), vd -0.100.", None, None, False),
    ("C7", "C7:E7", "F7", "h sàn (mm):", "150 = sàn 2 bên (theo Y2:Y3)\nMã DCE: 150/0 không tai · 150/1 tai trái · 150/2 tai phải · _150 sàn lật",
     None, "Chiều dày sàn (mm).\nChỉ nhập số: phía sàn / sàn lật lấy theo ô Y2, Y3.\nNhập mã DCE thì bỏ qua Y2:Y3:\n150 = 2 tai sàn · 150/0 = không tai · 150/1 = tai trái · 150/2 = tai phải\n_150 = sàn lật (dầm úp) · _150/1 ...", False),
    ("C8", "C8:E8", "F8", "Dầm ngắn L(m) <", "Nhịp có Ltt nhỏ hơn giá trị này (m): đai rải đều 1 bước.", None, None, False),
    ("G2", "G2:I2", "J2", "Thép gia cường gối:", "a|b|c: a = vươn mũ gối xL, b = vùng đai dày xL (hoặc mm nếu >1), c = TC trên bụng nhịp.\n/ = từ mép cột · \\ = từ tim cột · | = L tim-tim, đo từ mép.\nTrống = QS_DAMSET", "KGOI",
     "HỆ SỐ THÉP GIA CƯỜNG GỐI (J2)\nDấu phân cách quyết định cách đo:\n  0.3/0.25/0.15  : L thông thủy, tính từ MÉP cột\n  0.3\\0.25\\0.15 : L tim-tim, tính từ TIM cột\n  0.3|0.25|0.15  : L tim-tim, tính từ MÉP cột\na = vươn thép mũ gối (x L), có thể theo lớp: 0.3;0.25\nb = vùng đai dày (x L; > 1 hiểu là mm, vd 0.3/600/0.15)\nc = thép tăng cường TRÊN bụng nhịp (L - 2cL)\nThêm +15d: 0.3+15d\nĐể trống = lấy QS_DAMSET.", False),
    ("G3", "G3:I3", "J3", "Thép gia cường nhịp:", "a\\b\\c: a = TC dưới bụng nhịp cách mép aL, b = TC dưới tại gối vươn bL, c = nhịp đầu/cuối.\nTrống = QS_DAMSET", "KNHIP",
     "HỆ SỐ THÉP GIA CƯỜNG NHỊP (J3), cùng quy ước dấu / \\ | như J2\na = thép TC dưới bụng nhịp: L - 2aL\nb = thép TC dưới tại gối: vươn bL\nc = nhịp đầu / cuối: đầu thanh TC dưới cách mép gối biên cL\nĐể trống = lấy QS_DAMSET.", False),
    ("G4", "G4:I4", "J4", "Neo thép trên (d):", "40 = 40d · 10;12-30d/40d · 10-500mm/12-600mm/...\nTrống = QS_DAMSET", "NEO", None, False),
    ("G5", "G5:I5", "J5", "Neo thép dưới (d):", "40 = 40d · 10;12-30d/40d · 10-350mm/12-420mm/...\nTrống = QS_DAMSET", "NEO", None, False),
    ("G6", "G6:I6", "J6", "Neo thép giá (d):", "40 = 40d · 10;12-30d/40d · 10-350mm/...\nTrống = QS_DAMSET", "NEO", None, False),
    ("G7", "G7:I7", "J7", "Tên trục MCN dầm:", "Tên trục ghi ở mặt cắt ngang (lưu, chưa dùng khi vẽ).", None, None, False),
    ("G8", "G8:I8", "J8", "Lệch trục / Ẩn MC dọc:", "Ltt = thêm dim nhịp thông thủy.\n(100, Y, 100/Y: lưu theo DCE, chưa dùng.)", "LECHMCN", None, False),
    ("K2", "K2:M2", "N2", "ĐK thép C đỡ lớp TC:", "Ø thép C đỡ lớp tăng cường (mm). Trống = Ø đai.", None, None, False),
    ("K3", "K3:M3", "N3", "KC thép C đỡ lớp TC:", ">0: tự vẽ khi có ≥ 2 lớp · <0 (-1500): luôn vẽ khi > 2 thanh · 1500mm: luôn vẽ khi ≥ 2 thanh · 0/trống: không vẽ", "CDOKC",
     "KHOẢNG RẢI THÉP C ĐỠ LỚP TĂNG CƯỜNG (mm)\n> 0 : tự phân biệt lúc cần vẽ (có ≥ 2 lớp)\n= 0 : không vẽ\n< 0 : luôn vẽ khi số thép dọc > 2\n1500mm : luôn vẽ khi số thép dọc ≥ 2", False),
    ("K4", "K4:M4", "N4", "Đai C nối giữa thép giá:", "None / So le / So le 2 / Toàn bộ", "DAIC", None, True),
    ("K5", "K5:M5", "N5", "ĐK đai C nối 2 thép giá:", "Ø đai C nối thép giá (mm). Trống = Ø đai.", None, None, False),
    ("K6", "K6:M6", "N6", "KC đai C nối 2 thép giá:", "400 = móc đai C quay LÊN · 400mm = quay XUỐNG\nTrống = theo bước đai chính.", "KCDAIC", None, False),
    ("K7", "K7:M7", "N7", "Giật cao độ dầm:", "Mép trên / Mép dưới / Tùy chỉnh (lưu, chưa dùng khi vẽ).", "GIATCD", None, True),
    ("K8", "K8:M8", "N8", "Dày BT lót dưới dầm:", "100 = BT lót dày 100 · 100/50 = dày 100, mở rộng mỗi bên 50 (mm).", None, None, False),
]

CHO = [(3, "Kiểu thép chờ", "CHOKIEU", True,
        "Khong = không chờ\nCho thang = kéo thẳng ra ngoài đầu dầm L chờ, không bẻ ke\nCoupler = dừng tại mép, vẽ ký hiệu coupler"),
       (4, "Lớp được chờ", "CHOLOP", False, "T = trên · B = dưới · G = giá. Ghép: TBG, TB ..."),
       (5, "Thanh được chờ", "CHOTHEP", True, "Tat ca = mọi thanh tới đầu dầm\nChay suot = chỉ thép chạy suốt"),
       (6, "L chờ", "CHOL", False, "Chờ thẳng: AUTO = L nối ngoài vùng · 40d · 1200 (mm)\nCoupler: đoạn ra ngoài mép 0 · 100 · 100/300 = so le (½ thanh 100, ½ thanh 300); AUTO = QS_DAMSET"),
       (7, "Chờ so le 50%", "KHONGCO", True, "Co = 1/2 số thanh chờ L, 1/2 chờ 2L + KC mối nối"),
       (8, "Tên dầm chờ nối", None, False, "Tên dầm / zone sau mà thép chờ (thẳng hoặc coupler) nối sang, vd B138 – ghi lên MC dọc và shop: 'COUPLER NỐI DẦM B138'. Để trống = không ghi.")]


def as_area_expr(n):
    """Dien tich (cm2) cua chuoi thep da chuan hoa o o n (toi da 2 nhom a+b)."""
    g1 = f'LEFT({n},FIND("+",{n}&"+")-1)'
    g2 = f'MID({n},FIND("+",{n}&"+")+1,99)'
    def G(x): return f'IFERROR(LEFT({x},FIND("t",{x})-1)*MID({x},FIND("t",{x})+1,2)^2*PI()/400,0)'
    return f"{G(g1)}+{G(g2)}"


def norm_expr(x):
    """chuan hoa ky hieu thep: 3T28 / 3%%c28 / 3Ø28 / 3f28 / 3d28 -> 3t28"""
    e = f'LOWER({x})'
    for a in (" ", "%%c", "ø", "φ", "f", "d"):
        e = f'SUBSTITUTE({e},"{a}","{"" if a == " " else "t"}")'
    return e


def side_expr(x, side):
    if side == "L":
        return f'IF(ISNUMBER(FIND(";",{x})),LEFT({x},FIND(";",{x})-1),{x})'
    return f'IF(ISNUMBER(FIND(";",{x})),MID({x},FIND(";",{x})+1,99),{x})'


def num_expr(x):
    """so tu o co the la chuoi '-149.2' (chiu duoc dau thap phan , hoac .)"""
    return f'IFERROR({x}*1,IFERROR(SUBSTITUTE({x},".",",")*1,IFERROR(SUBSTITUTE({x},",",".")*1,0)))'


def build_sheet(wb, title, example):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = "FFC000" if example else "92D050"
    ws.sheet_view.zoomScale = 90
    ws.sheet_view.showGridLines = False
    widths = {"A": 30, "B": 9}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    for c in GCOLS: ws.column_dimensions[CL(c)].width = 9.5
    for c in NCOLS: ws.column_dimensions[CL(c)].width = 10.5
    for c in range(CLAST + 1, CLAST + 4): ws.column_dimensions[CL(c)].width = 3

    dvs = {}
    def dv(key_prompt, cells, title_, prompt, lst=None, strict=False):
        k = (key_prompt, lst, strict)
        if k not in dvs:
            d = DataValidation(type="list" if lst else None, formula1=list_ref(lst) if lst else None,
                               allow_blank=True, showDropDown=False)
            d.showInputMessage = True
            d.promptTitle = title_[:32]
            d.prompt = prompt[:255]
            d.showErrorMessage = bool(strict)
            if strict:
                d.errorTitle = "QS_DAM"; d.error = "Chọn giá trị trong danh sách."
            ws.add_data_validation(d)
            dvs[k] = d
        for c in cells: dvs[k].add(c)

    def put(addr, v, font=None, fill=None, align=CEN, border=B_ALL):
        c = ws[addr]
        if v is not None: c.value = v
        if font: c.font = font
        if fill: c.fill = FILL(fill)
        if align: c.alignment = align
        if border: c.border = border
        return c

    # ---------------- logo / marker / title ----------------
    ws.merge_cells("A1:B3")
    put("A1", "QS_DAM_V2", F(size=18, bold=True, color="C00000"), "FFF2CC")
    ws["A1"].comment = Comment("Mã nhận dạng sheet nhập liệu QS_DAM V2 – KHÔNG xóa / sửa ô này.\n"
                               "Bố cục ô giống sheet DCE_Pro_Beam (DAM.xlsm).", "QS_DAM", width=320, height=90)
    ws.merge_cells("A4:B4")
    put("A4", "NHẬP LIỆU DẦM – QS_VEDAM / QS_SHOPDAM", F(bold=True, size=9), "FFF2CC")
    ws.merge_cells("A5:B7")
    put("A5", "Nguyễn Xuân Phát\nbanhbaonxp@gmail.com\nQS_DAM v3.1", F(bold=True, color="0000CC", size=9), C_AUTH)
    ws.merge_cells("A8:B8")
    put("A8", "Ô trắng = nhập · ▼ = có danh sách", F(italic=True, size=8, color="595959"), "FFFFFF")
    ws.merge_cells("C1:N1")
    put("C1", "THÔNG TIN CẤU KIỆN", F(bold=True, size=12), C_TITLE)
    for c in range(4, 15): ws.cell(1, c).border = B_ALL

    # ---------------- header ----------------
    for lab_cell, mrg, val_cell, text, prompt, lst, helptxt, strict in HEAD:
        ws.merge_cells(mrg)
        put(lab_cell, text, F(bold=True), C_LABEL, NOWRAP)
        for c in ws[mrg][0]: c.border = B_ALL; c.fill = FILL(C_LABEL)
        put(val_cell, EX_HEAD.get(val_cell) if example else None, F(bold=True, color="0000FF"), "FFFFFF", NOWRAP)
        dv(val_cell, [val_cell], text.rstrip(":"), prompt, lst, strict)
        if helptxt:
            ws[lab_cell].comment = Comment(helptxt, "QS_DAM", width=380, height=200)
    ws["F6"].number_format = "+0.000;-0.000;±0.000"

    # ---------------- thep cho 2 dau + tuy chon san (phan mo rong QS) ----------------
    ws.merge_cells("P1:T1")
    put("P1", "THÉP CHỜ 2 ĐẦU DẦM", F(bold=True, size=11), C_TITLE)
    for c in range(17, 21): ws.cell(1, c).border = B_ALL
    ws.merge_cells("P2:R2")
    put("P2", "Mạch ngừng / zone sau", F(bold=True, size=9), C_LABEL)
    for c in range(17, 19): ws.cell(2, c).border = B_ALL
    put("S2", "TRÁI", F(bold=True), C_GOIHDR); put("T2", "PHẢI", F(bold=True), C_GOIHDR)
    for r, lab, lst, strict, prm in CHO:
        ws.merge_cells(f"P{r}:R{r}")
        put(f"P{r}", lab, F(bold=True, size=9), C_LABEL, NOWRAP)
        for c in range(17, 19): ws.cell(r, c).border = B_ALL
        for col in "ST":
            put(f"{col}{r}", EX_HEAD.get(f"{col}{r}") if example else None, F(bold=True, color="0000FF"), "FFFFFF")
        dv(f"CHO{r}", [f"S{r}", f"T{r}"], lab, prm, lst, strict)
    ws["P3"].comment = Comment("THÉP CHỜ 2 ĐẦU DẦM (khi đầu dầm là mạch ngừng, chờ sang zone / đợt đổ sau)\n"
                               "• Cho thang: thép kéo thẳng ra ngoài đầu dầm đoạn L chờ, không bẻ ke.\n"
                               "• Tên dầm chờ nối (S8 / T8): ghi 'THÉP CHỜ / COUPLER NỐI DẦM …' trên MC dọc, shop và tên dầm trong vùng nét khuất.\n"
                               "• Coupler: thép dừng tại mép (L chờ = 0) hoặc kéo ra ngoài mép: 100, hoặc so le 100/300 (½ số thanh 100, ½ số thanh 300).\n"
                               "• So le: 1/2 số thanh chờ L, 1/2 chờ 2L + KC mối nối.\n"
                               "Bản vẽ có nét khuất dầm zone sau, đường MẠCH NGỪNG, dim L chờ; QS_SHOPDAM cắt đúng chiều dài.\n"
                               "Để trống = theo QS_DAMSET.", "QS_DAM", width=420, height=170)
    ws.merge_cells("V1:Y1")
    put("V1", "TÙY CHỌN SÀN / ĐAI CON", F(bold=True, size=11), C_TITLE)
    for c in range(23, 26): ws.cell(1, c).border = B_ALL
    for r, lab, lst, prm in ((2, "Tai sàn (phía có sàn)", "TAISAN", "2 ben / Trai / Phai / Khong\nChỉ dùng khi F7 là số."),
                             (3, "Sàn lật (dầm úp)", "KHONGCO", "Co = sàn nằm ở đáy dầm. Chỉ dùng khi F7 là số.")):
        ws.merge_cells(f"V{r}:X{r}")
        put(f"V{r}", lab, F(bold=True, size=9), C_LABEL, NOWRAP)
        for c in range(23, 25): ws.cell(r, c).border = B_ALL
        put(f"Y{r}", EX_HEAD.get(f"Y{r}") if example else None, F(bold=True, color="0000FF"), "FFFFFF")
        dv(f"SAN{r}", [f"Y{r}"], lab, prm, lst, True)
    # dai con: nhanh dai trong theo vi tri thanh lop 1 tren
    for r, lab, lst, strict, prm in (
            (4, "Đai con đi hết dầm", "KHONGCO", True,
             "Co = 1 khai báo (Y5) cho cả dầm.\nKhong = vùng gối (vùng đai dày 2 đầu nhịp) dùng Y5, vùng nhịp dùng Y6.\nTrống = theo QS_DAMSET trang 2."),
            (5, "Đai con toàn dầm / vùng gối", "DAICON", False,
             "Theo số thứ tự thanh lớp 1 trên (trái → phải):\n3 = đai C tại thanh 3\n2_4 = đai Q (kín) ôm thanh 2..4\n2-4 = đai U thanh 2..4\nGhép: 3,2_4 · 0 = không có · AUTO/trống = tự động"),
            (6, "Đai con vùng nhịp", "DAICON", False,
             "Dùng khi Y4 = Khong. Cú pháp như Y5: 2 · 2_4 · 2-4 · 3,2_4 · 0 · AUTO")):
        ws.merge_cells(f"V{r}:X{r}")
        put(f"V{r}", lab, F(bold=True, size=9), C_LABEL, NOWRAP)
        for c in range(23, 25): ws.cell(r, c).border = B_ALL
        put(f"Y{r}", EX_HEAD.get(f"Y{r}") if example else None, F(bold=True, color="0000FF"), "FFFFFF")
        dv(f"CON{r}", [f"Y{r}"], lab, prm, lst, strict)
        if r >= 5: ws[f"Y{r}"].number_format = "@"      # dang Text: "2-4" khong bi Excel doi thanh ngay thang
    ws["V4"].comment = Comment(
        "ĐAI CON (nhánh đai bên trong) theo vị trí thanh thép lớp 1 trên, đánh số 1..n từ trái sang phải:\n"
        "• 2 → đai C (1 nhánh) tại thanh số 2\n• 2_4 → đai Q (đai kín) ôm từ thanh 2 đến thanh 4\n"
        "• 2-4 → đai U từ thanh 2 đến thanh 4\n• ghép nhiều loại: 3,2_4 (đai C thanh 3 + đai Q thanh 2..4)\n"
        "• 0 = không có đai con ; AUTO / trống = tự động (QS_DAMSET trang 2)\n"
        "Vùng gối có gia cường nên số thanh khác vùng nhịp → Y4 = Khong để khai báo riêng vùng gối (Y5) và vùng nhịp (Y6).\n"
        "Chỉ số vượt số thanh thực tế được tự thu về thanh cuối.", "QS_DAM", width=440, height=190)
    for c in ("P", "Q", "R", "S", "T", "V", "W", "X", "Y"):
        ws.column_dimensions[c].width = max(ws.column_dimensions[c].width or 0, 10.5 if c in "ST Y" else 9)

    for r in range(1, 9): ws.row_dimensions[r].height = 17
    # ---------------- dong 9: tinh trang so lieu ----------------
    put("A9", "TÌNH TRẠNG SỐ LIỆU:", F(bold=True, color="FFFFFF"), "404040", LEFT)
    put("B9", None, None, "404040")
    ws.merge_cells(f"C9:{LAST}9")
    for c in range(C0, CLAST + 1): ws.cell(9, c).fill = FILL("404040"); ws.cell(9, c).border = B_ALL
    ws["C9"].font = F(bold=True, color="FFFFFF", size=11)
    ws["C9"].alignment = LEFT

    # ---------------- dong 10: tieu de goi / nhip ----------------
    put("A10", "HẠNG MỤC", F(bold=True), C_GOIHDR)
    put("B10", "Chạy suốt", F(bold=True, size=9), C_GOIHDR)
    for i, c in enumerate(GCOLS):
        put(f"{CL(c)}10", f"Gối {i + 1}", F(bold=True), C_GOIHDR)
    for i, c in enumerate(NCOLS):
        put(f"{CL(c)}10", f"Nhịp {i + 1}", F(bold=True, color="C00000"), C_NHIPHDR)
    ws.row_dimensions[10].height = 18

    # ---------------- dong 11..30 ----------------
    ws.merge_cells("A13:A17"); ws.merge_cells("A18:A22")
    put("A13", "THÉP TĂNG CƯỜNG\nLỚP TRÊN", F(bold=True), C_LABEL)
    put("A18", "THÉP TĂNG CƯỜNG\nLỚP DƯỚI", F(bold=True), C_LABEL)
    for r in range(13, 23): ws.cell(r, 1).border = B_ALL; ws.cell(r, 1).fill = FILL(C_LABEL)
    ws["A13"].comment = Comment(HELP_ROW[13], "QS_DAM", width=420, height=130)
    ws["A18"].comment = Comment(HELP_ROW[18], "QS_DAM", width=420, height=110)
    for r in range(11, 31):
        lab, sub, gp, np_ = ROWS[r]
        if lab:
            if r >= 23:
                ws.merge_cells(f"A{r}:B{r}")
                put(f"A{r}", lab, F(bold=True, size=9), "FFFFFF", LEFT)
                ws.cell(r, 2).border = B_ALL
            else:
                put(f"A{r}", lab, F(bold=True, size=9), "FFFFFF", LEFT)
            if r in HELP_ROW:
                ws[f"A{r}"].comment = Comment(HELP_ROW[r], "QS_DAM", width=420, height=150)
        if sub:
            put(f"B{r}", sub, F(bold=True, size=9), C_LABEL)
        ws.row_dimensions[r].height = 27 if r in (11, 12) else (21 if r >= 23 else 16)
        for c in range(C0, CLAST + 1):
            goi = c in GCOLS
            addr = f"{CL(c)}{r}"
            if 13 <= r <= 22:
                put(addr, None, F(size=9), C_LAYERG if goi else C_LABEL, CEN, B_HAIR)
            elif r >= 23 and goi:
                if r == 30:
                    put(addr, "", F(size=9, color="FFFFFF"), C_NA)
                else:
                    put(addr, None, F(bold=True, size=9, color="FFFFFF"), C_GOIDARK)
            elif r == 11:
                put(addr, None, F(bold=True, color="000000" if goi else "C00000"), "F2F2F2" if goi else "FFFFFF")
            else:
                put(addr, None, F(size=9), "F2F2F2" if goi else "FFFFFF")
        gcells = [f"{CL(c)}{r}" for c in GCOLS]
        ncells = [f"{CL(c)}{r}" for c in NCOLS]
        if gp: dv(f"G{r}", gcells, gp[0], gp[1], gp[2])
        if np_: dv(f"N{r}", ncells, np_[0], np_[1], np_[2])
    # B11 / B12: thep chay suot
    for r, t in ((11, "TRÊN"), (12, "DƯỚI")):
        put(f"B{r}", EX_HEAD.get(f"B{r}") if example else None, F(bold=True, size=11, color="0000FF"), "FFFFFF")
        dv(f"B{r}", [f"B{r}"], f"Thép chạy suốt {t}", f"Thép chạy suốt {t}, vd 3t28 · 2t28+1t25.\nĐổi thép theo nhịp: dòng {23 if r == 11 else 25}.", "THEP")
    for r in (11, 12):
        ws[f"A{r}"].font = F(bold=True, size=9)

    # du lieu vi du
    if example:
        for r, vals in EX_GRID.items():
            for i, v in enumerate(vals):
                if v is not None: ws.cell(r, C0 + i).value = v
        for r, vals in EX_AS.items():
            for i, v in enumerate(vals):
                if v is not None: ws.cell(r, C0 + i).value = v

    # vien dam bao quanh luoi
    for r in range(10, 31):
        for c in (1, CLAST):
            cell = ws.cell(r, c)
            b = cell.border
            cell.border = Border(left=med if c == 1 else b.left, right=med if c == CLAST else b.right,
                                 top=b.top, bottom=b.bottom)
    for c in range(1, CLAST + 1):
        for r, side in ((10, "top"), (30, "bottom")):
            cell = ws.cell(r, c); b = cell.border
            cell.border = Border(left=b.left, right=b.right, top=med if side == "top" else b.top,
                                 bottom=med if side == "bottom" else b.bottom)

    # ================= HANG AN: tinh toan phu tro (dong 60+) =================
    H = {}
    r0 = 60
    names = ["W", "X", "AX", "ACT", "SHOW", "HASL", "HASR", "LECH"]
    layers = [("T0", None)] + [(f"T{k}", 12 + k) for k in range(1, 6)] + [("B0", None)] + [(f"B{k}", 23 - k) for k in range(1, 6)]
    for nm, _ in layers:
        for sd in "LR":
            names += [f"N_{nm}{sd}", f"A_{nm}{sd}"]
    names += ["ASTL", "ASTR", "ASBL", "ASBR", "YTL", "YTR", "YBL", "YBR", "BAD", "NB11", "NB12"]
    for i, nm in enumerate(names):
        H[nm] = r0 + i
        ws.cell(r0 + i, 1).value = f"(phụ trợ) {nm}"
        ws.cell(r0 + i, 1).font = F(size=8, color="808080")
        ws.row_dimensions[r0 + i].hidden = True
        ws.row_dimensions[r0 + i].outlineLevel = 1
    ws.cell(r0 - 1, 1).value = "▼ Hàng ẩn phụ trợ (dòng 60+) – dùng cho kiểm tra, không nhập."
    ws.cell(r0 - 1, 1).font = F(size=8, italic=True, color="808080")
    ws.sheet_format.outlineLevelRow = 1

    def h(nm, col): return f"{CL(col)}{H[nm]}"
    ws[f"B{H['NB11']}"] = "=" + norm_expr("$B$11")
    ws[f"B{H['NB12']}"] = "=" + norm_expr("$B$12")
    nb = {"T": f"$B${H['NB11']}", "B": f"$B${H['NB12']}"}
    for c in range(C0, CLAST + 1):
        L = CL(c); goi = c in GCOLS
        prev, nxt = CL(c - 1), CL(c + 1)
        if goi:
            ws[h("W", c)] = f'=IF({L}11="","",IFERROR(LEFT({L}11,MIN(FIND({{"x","/"}},{L}11&"x/"))-1)*1,0))'
            ws[h("LECH", c)] = "=" + num_expr(f"{L}27")
        else:
            ws[h("W", c)] = f'=IF(ISNUMBER({L}11),{L}11,IFERROR({L}11*1,""))'
        ws[h("X", c)] = "=0" if c == C0 else f"={h('X', c - 1)}+N({h('W', c - 1)})"
        if goi:
            left_act = f"N({h('ACT', c - 1)})" if c > C0 else "0"
            right_act = f"N({h('ACT', c + 1)})" if c < CLAST else "0"
            ws[h("ACT", c)] = f"=IF({left_act}+{right_act}>0,1,0)"
            ws[h("HASL", c)] = f"={left_act}"
            ws[h("HASR", c)] = f"={right_act}"
            ws[h("AX", c)] = f'=IF({h("ACT", c)}=1,{h("X", c)}+N({h("W", c)})/2+{h("LECH", c)},"")'
            ls = f"N({h('SHOW', c - 1)})" if c > C0 else "0"
            rs = f"N({h('SHOW', c + 1)})" if c < CLAST else "0"
            ws[h("SHOW", c)] = f"=IF({ls}+{rs}>0,1,0)"
        else:
            ws[h("ACT", c)] = f"=IF(N({h('W', c)})>0,1,0)"
            ws[h("HASL", c)] = "=1"; ws[h("HASR", c)] = "=1"
            prevn = f"N({h('ACT', c - 2)})" if c > C0 + 1 else "1"
            ws[h("SHOW", c)] = f"=IF({h('ACT', c)}+{prevn}>0,1,0)"
        # thep chay suot hieu dung (doi thep o dong 23 / 25 cot nhip)
        for tp, row in (("T", 23), ("B", 25)):
            for sd in "LR":
                end = (c - 1 if sd == "L" else c + 1) if goi else c
                cell = h(f"N_{tp}0{sd}", c)
                if end < C0 + 1:
                    ws[cell] = f"={nb[tp]}"
                    continue
                rg = f"${CL(C0 + 1)}${row}:{CL(end)}${row}"
                lk = f'LOOKUP(2,1/ISNUMBER(FIND(";",{rg})),{rg})'
                part = f'MID({lk},FIND(";",{lk})+1,99)'
                ws[cell] = (f'=IFERROR(IF(ISNUMBER({part}*1),LEFT({nb[tp]},FIND("t",{nb[tp]}))&{part},'
                            f'{norm_expr(part)}),{nb[tp]})')
        for nm, src in layers:
            for sd in "LR":
                if src is not None:
                    ws[h(f"N_{nm}{sd}", c)] = "=" + norm_expr(side_expr(f"{L}{src}", sd))
                ws[h(f"A_{nm}{sd}", c)] = "=" + as_area_expr(h(f"N_{nm}{sd}", c))
        for tp in "TB":
            for sd in "LR":
                ws[h(f"AS{tp}{sd}", c)] = "=" + "+".join(h(f"A_{tp}{k}{sd}", c) for k in range(6))
        # As yeu cau "a" hoac "a/b" (L/R)
        for tp, row in (("T", 37), ("B", 40)):
            y = f"{L}{row}"
            left = f'LEFT({y},FIND("/",{y}&"/")-1)'
            ws[h(f"Y{tp}L", c)] = f'=IF({y}="","",{num_expr(left)})'
            right = 'MID(%s,FIND("/",%s)+1,20)' % (y, y)
            ws[h(f"Y{tp}R", c)] = f'=IF({y}="","",IF(ISNUMBER(FIND("/",{y})),{num_expr(right)},{h(f"Y{tp}L", c)}))'
        bad = "+".join(f'(({L}{r}<>"")*({L}{r}<>"-")*({h(f"A_{nm}L", c)}+{h(f"A_{nm}R", c)}=0))'
                       for nm, r in layers if r is not None)
        ws[h("BAD", c)] = "=" + bad

    # ================= DONG 32..43: KIEM TRA =================
    def section(r, text):
        ws.merge_cells(f"A{r}:{LAST}{r}")
        put(f"A{r}", text, F(bold=True), C_SEC, LEFT)
        for c in range(2, CLAST + 1): ws.cell(r, c).fill = FILL(C_SEC); ws.cell(r, c).border = B_ALL
        ws.row_dimensions[r].height = 18

    section(32, "KIỂM TRA HÌNH HỌC (tự tính – đối chiếu mặt bằng kết cấu)")
    calc_rows = {
        33: ("Gối: tọa độ trục (gốc = trục gối 1)  ·  Nhịp: KC trục-trục (mm)", True),
        34: ("Nhịp: KC trục theo mặt bằng (nhập để đối chiếu)", False),
        35: ("Nhịp: Ltt suy ra từ KC trục ở dòng 34 (mm)", True),
    }
    section(36, "KIỂM TRA DIỆN TÍCH THÉP (cm²) – thép chạy suốt hiệu dụng + tăng cường lớp 1..5 (gối: trái/phải)")
    calc_rows.update({
        37: ("As yêu cầu TRÊN (nhập; gối: trái/phải vd 2.63/1.45)", False),
        38: ("As bố trí TRÊN", True),
        39: ("Hệ số As bố trí / yêu cầu TRÊN", True),
        40: ("As yêu cầu DƯỚI (nhập)", False),
        41: ("As bố trí DƯỚI", True),
        42: ("Hệ số As bố trí / yêu cầu DƯỚI", True),
    })
    section(43, "KIỂM TRA SỐ LIỆU")
    calc_rows[44] = ("Kết quả kiểm tra từng cột", True)
    for r, (lab, calc) in calc_rows.items():
        ws.merge_cells(f"A{r}:B{r}")
        put(f"A{r}", lab, F(bold=True, size=9), C_CALC if calc else C_LABEL, LEFT)
        ws.cell(r, 2).border = B_ALL
        ws.row_dimensions[r].height = 27
        for c in range(C0, CLAST + 1):
            put(f"{CL(c)}{r}", None, F(size=9, italic=calc, color="404040" if calc else "0000FF"),
                C_CALC if calc else "FFFFFF")

    for c in range(C0, CLAST + 1):
        L = CL(c); goi = c in GCOLS; A = lambda nm: h(nm, c)
        if goi:
            ws[f"{L}33"] = f'=IF({A("ACT")}=1,{A("AX")}-$C${H["AX"]},"")'
            ws[f"{L}34"].fill = FILL(C_CALC); ws[f"{L}35"].fill = FILL(C_CALC)
        else:
            ws[f"{L}33"] = f'=IF({A("ACT")}=1,IFERROR({h("AX", c + 1)}-{h("AX", c - 1)},""),"")'
            ws[f"{L}35"] = (f'=IF({L}34="","",{L}34-(N({h("W", c - 1)})/2-{h("LECH", c - 1)})'
                            f'-(N({h("W", c + 1)})/2+{h("LECH", c + 1)}))')
        for r in (33, 34, 35): ws[f"{L}{r}"].number_format = "#,##0.0;-#,##0.0;0"
        for tp, rb, rr in (("T", 38, 39), ("B", 41, 42)):
            tl, tr = A(f"AS{tp}L"), A(f"AS{tp}R")
            if goi:
                ws[f"{L}{rb}"] = (f'=IF({A("ACT")}=1,IF({A("HASL")}*{A("HASR")}=1,FIXED({tl},2,TRUE)&"/"&FIXED({tr},2,TRUE),'
                                  f'IF({A("HASL")}=1,FIXED({tl},2,TRUE),FIXED({tr},2,TRUE))),"")')
            else:
                ws[f"{L}{rb}"] = f'=IF({A("ACT")}=1,ROUND({tl},2),"")'
                ws[f"{L}{rb}"].number_format = "0.00"
            yl, yr = A(f"Y{tp}L"), A(f"Y{tp}R")
            rl = f'IF(AND({A("HASL")}=1,N({yl})>0),{tl}/{yl},99)'
            rr_ = f'IF(AND({A("HASR")}=1,N({yr})>0),{tr}/{yr},99)'
            ws[f"{L}{rr}"] = f'=IF(OR({A("ACT")}=0,{yl}=""),"",IF(MIN({rl},{rr_})=99,"",ROUND(MIN({rl},{rr_}),2)))'
            ws[f"{L}{rr}"].number_format = "0.00"
        if goi:
            ws[f"{L}44"] = (f'=IF({A("ACT")}=0,"",IF({A("W")}="","Thiếu KT gối",IF({L}26="","Thiếu tên trục",'
                            f'IF({A("BAD")}>0,"Sai KH thép","OK"))))')
        else:
            gap = f'AND({CL(c - 2)}${H["ACT"]}=0,{A("ACT")}=1)' if c > C0 + 1 else "FALSE"
            ws[f"{L}44"] = (f'=IF({A("ACT")}=0,IF({L}11="","","Sai Ltt"),IF({gap},"Nhịp đứt quãng",'
                            f'IF({L}26="","Thiếu đai",IF({A("BAD")}>0,"Sai KH thép","OK"))))')
        ws[f"{L}44"].font = F(bold=True, size=9)

    # dong 9: tong hop
    rngACT = f"{FIRST}{H['ACT']}:{LAST}{H['ACT']}"
    nspan = f"SUMPRODUCT((MOD(COLUMN({rngACT}),2)=0)*{rngACT})"
    total = f"SUMIF({rngACT},1,{FIRST}{H['W']}:{LAST}{H['W']})"
    a11 = as_area_expr("$B$%d" % H["NB11"]); a12 = as_area_expr("$B$%d" % H["NB12"])
    nerr = (f'SUMPRODUCT(({FIRST}44:{LAST}44<>"")*({FIRST}44:{LAST}44<>"OK"))'
            f'+(({a11})=0)+(({a12})=0)')
    ws["C9"] = (f'=IF({nspan}=0,"Chưa có nhịp: nhập bề rộng gối + L thông thủy ở dòng 11",'
                f'IF({nerr}=0,"✔  Số liệu hợp lệ – "&{nspan}&" nhịp, tổng chiều dài dầm L = "&FIXED({total},0)&" mm",'
                f'"✖  Có "&({nerr})&" lỗi – xem dòng 44 (KIỂM TRA SỐ LIỆU) và ô B11 / B12"))')

    # ================= DINH DANG CO DIEU KIEN =================
    grid = f"{FIRST}11:{LAST}30"
    # cot khong dung (ngoai so nhip) -> xam; chua lai 1 nhip trong ke tiep de nhap them
    inact = f'{FIRST}${H["SHOW"]}=0'
    ws.conditional_formatting.add(grid, FormulaRule(formula=[inact], fill=FILL("BFBFBF"), font=Font(color="7F7F7F"), stopIfTrue=True))
    ws.conditional_formatting.add(f"{FIRST}33:{LAST}44", FormulaRule(formula=[f'{FIRST}${H["ACT"]}=0'], fill=FILL("D9D9D9")))
    # o thep sai ky hieu -> do
    for r in range(13, 23):
        nm = next(n for n, s in layers if s == r)
        f = f'AND({FIRST}{r}<>"",{FIRST}{r}<>"-",{FIRST}${H[f"A_{nm}L"]}+{FIRST}${H[f"A_{nm}R"]}=0)'
        ws.conditional_formatting.add(f"{FIRST}{r}:{LAST}{r}", FormulaRule(formula=[f], fill=FILL("FF7C80"), font=Font(bold=True, color="9C0006")))
    for a, nm in (("B11", "NB11"), ("B12", "NB12")):
        ws.conditional_formatting.add(a, FormulaRule(formula=['(%s)=0' % as_area_expr("$B$%d" % H[nm])], fill=FILL("FF7C80")))
    # he so As
    for r in (39, 42):
        rg = f"{FIRST}{r}:{LAST}{r}"
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({FIRST}{r}),{FIRST}{r}<1)'], fill=FILL("FFC7CE"), font=Font(bold=True, color="9C0006")))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({FIRST}{r}),{FIRST}{r}>=1)'], fill=FILL("C6EFCE"), font=Font(bold=True, color="006100")))
    # Ltt suy ra lech Ltt nhap > 1 mm
    ws.conditional_formatting.add(f"{FIRST}35:{LAST}35", FormulaRule(
        formula=[f'AND(ISNUMBER({FIRST}35),ISNUMBER({FIRST}11),ABS({FIRST}35-{FIRST}11)>1)'], fill=FILL("FFEB9C"), font=Font(bold=True, color="9C5700")))
    # ket qua kiem tra
    rg = f"{FIRST}44:{LAST}44"
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'{FIRST}44="OK"'], fill=FILL("C6EFCE"), font=Font(bold=True, color="006100")))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND({FIRST}44<>"",{FIRST}44<>"OK")'], fill=FILL("FFC7CE"), font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add("C9", FormulaRule(formula=['LEFT($C$9,1)="✖"'], fill=FILL("C00000")))
    ws.conditional_formatting.add("C9", FormulaRule(formula=['LEFT($C$9,1)="✔"'], fill=FILL("00B050")))

    dv("AS37", [f"{CL(c)}37" for c in range(C0, CLAST + 1)], "As yêu cầu TRÊN (cm²)",
       "As yêu cầu thép TRÊN (cm²) từ tính toán, vd 2.34.\nGối: trái/phải vd 2.63/1.45.\nĐể trống = không kiểm tra.")
    dv("AS40", [f"{CL(c)}40" for c in range(C0, CLAST + 1)], "As yêu cầu DƯỚI (cm²)",
       "As yêu cầu thép DƯỚI (cm²), vd 2.46. Gối: trái/phải.\nĐể trống = không kiểm tra.")
    dv("KC34", [f"{CL(c)}34" for c in NCOLS], "KC trục mặt bằng (mm)",
       "Nhập khoảng cách trục-trục theo mặt bằng để suy ra Ltt (dòng 35).\nKhông ảnh hưởng bản vẽ.")

    ws.freeze_panes = "C11"
    ws.print_area = f"A1:{LAST}44"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_cols = "A:B"
    return ws


def build_list(wb):
    ws = wb.create_sheet("LIST")
    for k, vals in LISTS.items():
        col = LISTCOL[k]
        ws[f"{col}1"] = k
        ws[f"{col}1"].font = F(bold=True)
        for i, v in enumerate(vals):
            ws[f"{col}{i + 2}"] = v
        ws.column_dimensions[col].width = 14
    ws["A30"] = "Danh sách gợi ý / chọn cho sheet nhập liệu QS_DAM. Có thể thêm dòng (giữ trong vùng đã khai báo)."
    ws.sheet_state = "hidden"


def build_help(wb):
    ws = wb.create_sheet("HUONG_DAN", 0)
    ws.sheet_properties.tabColor = "4472C4"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    ws.merge_cells("A1:B1")
    ws["A1"] = "HƯỚNG DẪN NHẬP LIỆU QS_DAM V2 (bố cục giống sheet DCE_Pro_Beam)"
    ws["A1"].font = F(bold=True, size=14, color="FFFFFF"); ws["A1"].fill = FILL("4472C4")
    ws["A1"].alignment = CEN
    ws.row_dimensions[1].height = 26
    rows = [
        ("CÁCH DÙNG", "1. Mỗi dầm 1 sheet: chuột phải sheet MAU (hoặc sheet ví dụ) → Move or Copy → Create a copy, đổi tên theo dầm.\n"
                      "2. Nhập ô trắng. Chọn ô sẽ hiện gợi ý; ô có ▼ có danh sách chọn (vẫn gõ tay được, trừ ô bắt buộc chọn).\n"
                      "3. Xem dòng 9 (TÌNH TRẠNG SỐ LIỆU) và dòng 44 (KIỂM TRA) – sửa hết ô đỏ.\n"
                      "4. AutoCAD / ZWCAD: để sheet dầm cần vẽ ĐANG MỞ trong Excel → lệnh QS_VEDAM → nguồn Excel → chọn điểm mép TRÊN-TRÁI dầm.\n"
                      "5. QS_SHOPDAM: cắt thép / shop. QS_DAMSET, QS_DAMNOI: cài đặt móc đai, neo, chiều dài nối."),
        ("TƯƠNG THÍCH DCE", "Vị trí ô nhập giống hệt sheet DCE_Pro_Beam: F2:F8, J2:J8, N2:N8, B11, B12, lưới C11:AG30 (Gối 1 = cột C, Nhịp 1 = cột D ...).\n"
                            "→ Có thể copy nguyên vùng C2:N8 và B11:AG30 từ file DCE (DAM.xlsm) dán vào (Paste Values) là dùng được.\n"
                            "Ô A1 = QS_DAM_V2 là mã nhận dạng – không xóa. QS_DAM vẫn đọc được sheet DCE_Pro_Beam và sheet QS_DAM_V1 cũ."),
        ("PHẦN MỞ RỘNG QS", "• THÉP CHỜ 2 ĐẦU DẦM (S3:T7): kiểu (Khong / Cho thang / Coupler), lớp (TBG, TB, T, B), thanh (Tat ca / Chay suot), L chờ (AUTO / 40d / 1200; coupler 0 / 100 / 100/300 so le), tên dầm chờ nối (S8:T8), so le.\n"
                            "• TÙY CHỌN SÀN (Y2:Y3): tai sàn 2 ben / Trai / Phai / Khong, sàn lật – dùng khi F7 chỉ là số (nhập mã DCE 150/1, _150 thì bỏ qua Y2:Y3).\n"
                            "• Dòng 33-35: tọa độ trục, KC trục-trục tính từ bề rộng gối + Ltt + lệch trục; nhập KC trục mặt bằng (dòng 34) → Ltt suy ra (dòng 35, vàng nếu lệch Ltt đang nhập).\n"
                            "• Dòng 37-42: kiểm tra As (cm²) bố trí / yêu cầu – thép chạy suốt hiệu dụng (có tính đổi thép dòng 23/25) + tăng cường lớp 1..5; gối tách trái/phải.\n"
                            "• Dòng 44 + dòng 9: báo thiếu KT gối, thiếu tên trục, thiếu đai, Ltt sai, nhịp đứt quãng, ký hiệu thép sai (ô thép sai tô đỏ)."),
        ("THÉP CHỜ TRÊN BẢN VẼ", "• QS_VEDAM: nét khuất dầm zone sau, đường MẠCH NGỪNG, dim L chờ (chờ thẳng) hoặc ký hiệu COUPLER.\n"
                                 "• QS_SHOPDAM: thanh chờ thẳng được cắt đủ chiều dài có đoạn chờ; dải shop có đường MẠCH NGỪNG; đầu thanh coupler vẽ ký hiệu,\n"
                                 "  bảng thống kê thêm dòng COUPLER theo đường kính (số cái)."),
        ("BỐ CỤC SHOP", "Mặc định giống bản vẽ DCE: shop thép TRÊN phía trên MC dọc; dưới MC dọc là dải THÉP GIÁ sát trên dải THÉP DƯỚI;\n"
                        "shop THÉP ĐAI bên phải khung shop trên (1:1, dim từng đoạn, tag SH / L). Thay đổi ở QS_DAMSET → trang 5 (Bố cục shop)."),
        ("LẤY TỪ MẶT BẰNG KC", "Lệnh QS_DAMMB (AutoCAD): mở file này (có sheet MAU) → trên MBKC pick điểm ĐẦU và CUỐI dầm trên tim dầm (ngoài 2 gối biên).\n"
                               "Lệnh tự nhận dạng: bề rộng b (nét biên dầm), gối (cột / vách cắt tim; đầu dầm gối lên dầm → '400x700'), trục + lệch trục,\n"
                               "dầm phụ trong nhịp (vị trí từ tim gối, b x h từ text), dầm giao tại cột, tên + b x h dầm từ text 'B137 (400x1000)'\n"
                               "→ copy sheet MAU thành sheet mới tên dầm và ghi F2:F4, dòng 11, 24–29. Thép chủ, tăng cường… nhập tiếp bằng tay.\n"
                               "Trục: tùy chọn [Truc] của lệnh → chọn BLOCK trục (tên trục = thuộc tính / text trong block / text trong bóng trục) → pick 1 đường trục (layer trục)\n"
                               "→ pick 1 tên trục dạng text (layer tên trục, Enter nếu đã dùng block). [Mau] = thêm layer cột, dầm, text. Sửa ở QS_DAMSET trang 6.\n"
                               "Luôn kiểm tra lại số liệu trước khi vẽ."),
        ("ĐAI CON (Y4:Y6)", "Nhánh đai bên trong theo số thứ tự thanh lớp 1 trên (1..n từ trái sang phải): 2 = đai C tại thanh 2 ; 2_4 = đai Q (kín) ôm thanh 2..4 ;\n"
                            "2-4 = đai U thanh 2..4 ; ghép 3,2_4 ; 0 = không có ; AUTO / trống = tự động.\n"
                            "Y4 = Co: Y5 áp dụng cả dầm. Y4 = Khong: Y5 cho vùng gối (vùng đai dày 2 đầu nhịp), Y6 cho vùng nhịp. Trống = QS_DAMSET trang 2.\n"
                            "Đường kính / bước đai con = đai trong dòng 27 nếu có, không thì theo đai chính của vùng."),
        ("KÝ HIỆU THÉP", "3t28 = 3 thanh Ø28 (t, T, d, f, Ø, %%c đều được). 2t28+1t25 = nhiều loại. '-' = không có.\n"
                         "Gối: 3t28;5t28 = trái 3t28 / phải 5t28 ; ;2t28 = chỉ bên phải ; 2t25; = chỉ bên trái."),
        ("DÒNG 11 (B11 / lưới)", HELP_ROW[11]), ("DÒNG 12 (B12 / lưới)", HELP_ROW[12]),
        ("DÒNG 13-17", HELP_ROW[13]), ("DÒNG 18-22", HELP_ROW[18]),
        ("DÒNG 23", HELP_ROW[23]), ("DÒNG 24", HELP_ROW[24]), ("DÒNG 25", HELP_ROW[25]), ("DÒNG 26", HELP_ROW[26]),
        ("DÒNG 27", HELP_ROW[27]), ("DÒNG 28", HELP_ROW[28]), ("DÒNG 29", HELP_ROW[29]), ("DÒNG 30", HELP_ROW[30]),
    ]
    for lab, _m, cell, text, prompt, _l, helptxt, _s in HEAD:
        rows.append((f"{cell} – {text.rstrip(':')}", helptxt or prompt))
    rows += [
        ("CHƯA HỖ TRỢ", "Giật cấp mép trên / dưới (dòng 23, 25 số đầu ≠ 0), đổi tiết diện, móng băng (F3/F4/B12 dạng a/b/c), giật sàn T.LINK/-100/-200: "
                        "QS_DAM cảnh báo và bỏ qua / chỉ đổi thép. Các ô J7, N7, J8 (trừ Ltt) được lưu theo DCE nhưng chưa dùng khi vẽ."),
        ("GIÁ TRỊ MẶC ĐỊNH", "Neo, hệ số vùng, móc đai, chiều dài nối để trống = lấy theo QS_DAMSET. Các giá trị mặc định CHƯA XÁC NHẬN – "
                             "kỹ sư phụ trách kiểm tra theo thuyết minh dự án trước khi phát hành."),
    ]
    r = 3
    for a, b in rows:
        ws[f"A{r}"] = a; ws[f"B{r}"] = b
        ws[f"A{r}"].font = F(bold=True); ws[f"A{r}"].fill = FILL(C_LABEL)
        ws[f"A{r}"].alignment = LEFT; ws[f"B{r}"].alignment = LEFT
        ws[f"A{r}"].border = B_ALL; ws[f"B{r}"].border = B_ALL
        ws[f"B{r}"].font = F()
        ws.row_dimensions[r].height = max(18, 13.5 * sum(1 + len(x) // 125 for x in b.split("\n")) + 4)
        r += 1


def main():
    wb = Workbook()
    wb.remove(wb.active)
    build_help(wb)
    ex = build_sheet(wb, "VD_L1.B137", True)
    build_sheet(wb, "MAU", False)
    build_list(wb)
    wb.active = wb.sheetnames.index("VD_L1.B137")
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "VD_L1.B137"
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
